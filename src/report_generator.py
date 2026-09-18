import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import datetime

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill_color)
    tcPr.append(shd)

def generate_user_assessment_docx(
    detected_vars: dict,
    synthesis_summary: str,
    recommendations: list,
    output_path=None
):
    doc = Document()

    title = doc.add_heading(level=0)
    run_title = title.add_run("Darukaa.Earth: Ecological Assessment & Biodiversity Report")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(34, 84, 61)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(14)
    r_sub = p_sub.add_run(f"Generated on: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Lead AI Environmental Scientist")
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(74, 85, 104)

    doc.add_heading("1. Land Baseline Environmental Parameters", level=1)
    table = doc.add_table(rows=len(detected_vars) if detected_vars else 1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    if detected_vars:
        for idx, (k, v) in enumerate(detected_vars.items()):
            row = table.rows[idx]
            row.cells[0].text = k.replace("_", " ").title()
            row.cells[0].paragraphs[0].runs[0].font.bold = True
            row.cells[1].text = str(v)
            set_cell_background(row.cells[0], "F0FDF4")
    else:
        row = table.rows[0]
        row.cells[0].text = "Parameters"
        row.cells[1].text = "Standard environmental assessment query"

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    doc.add_heading("2. Ecosystem Diagnosis & Limiting Factors", level=1)
    doc.add_paragraph(synthesis_summary or "Multi-metric evaluation across soil, moisture, and vegetation layers.")

    doc.add_heading("3. Evidence-Backed Interventions & Quantitative Benchmarks", level=1)
    if not recommendations:
        doc.add_paragraph("No interventions generated yet. Run an assessment in the app to populate recommendations.")
    else:
        for idx, rec in enumerate(recommendations, start=1):
            h = doc.add_heading(f"{idx}. {rec.title}", level=2)
            h.paragraph_format.space_before = Pt(10)
            
            p_action = doc.add_paragraph()
            r_act_title = p_action.add_run("Targeted Action (What to do): ")
            r_act_title.font.bold = True
            p_action.add_run(rec.recommendation)

            p_mech = doc.add_paragraph()
            r_mech_title = p_mech.add_run("Scientific Mechanism (Why it works): ")
            r_mech_title.font.bold = True
            p_mech.add_run(rec.scientific_reasoning)

            p_chain = doc.add_paragraph()
            r_ch_title = p_chain.add_run("Multi-Metric Causal Chain: ")
            r_ch_title.font.bold = True
            p_chain.add_run(rec.multi_metric_chain)

            p_metrics = doc.add_paragraph()
            r_met_title = p_metrics.add_run("Quantified Metric Improvements:\n")
            r_met_title.font.bold = True
            for m in rec.impacted_metrics:
                p_m = doc.add_paragraph(f"? {m}", style="List Bullet")
                p_m.paragraph_format.space_after = Pt(2)

            p_cite = doc.add_paragraph()
            r_ci_title = p_cite.add_run("Scientific Citations: ")
            r_ci_title.font.bold = True
            p_cite.add_run("; ".join(rec.citations))

            p_meta = doc.add_paragraph()
            p_meta.add_run(f"Time Horizon: {rec.time_horizon} | Scientific Confidence: {rec.confidence_level}")
            p_meta.runs[0].font.italic = True
            p_meta.paragraph_format.space_after = Pt(14)

    if output_path is None:
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer
    else:
        doc.save(output_path)
        return output_path
